"""Transcribe the World Bank's historical income classification into a dbt seed.

    uv run --with openpyxl python -m scripts.build_income_classification_seed
    uv run --with openpyxl python -m scripts.build_income_classification_seed --xlsx X

**Why this exists.** `dim_country.income_group` comes from the World Bank's
`/country` endpoint, which publishes only the *current* answer — so the warehouse
stamps one classification onto every year it holds. The World Bank reclassifies
every July, and of the 163 economies classified in both 1987 and 2025, **74 (45%)
are in a different group at the two ends**. "Emissions by income group in 1990"
is therefore a 2026 grouping of 1990 emissions: the textbook Type-1 distortion,
and a wrong answer that looks entirely right.

`OGHIST.xlsx` is the fix. It is the same publisher, keyed on the same ISO3 codes
everything here already joins on, and it goes back to fiscal 1989.

**Why a seed and not a dlt resource**, which is the CBAM argument exactly (see
`scripts/build_cbam_seeds.py`): this is published once a year as a file drop with
a fiscal-year vintage, not served by the API the other World Bank resources read.
It is versioned by publication, not by scrape. dbt Labs' structure guide says not
to load source data with seeds, and this takes the same exception the CBAM seeds
take, for the same reason — the alternative is a dlt resource that re-fetches an
annual reference table on every run.

Three things this file cost, none of them guessable from the data:

* **The undated URL is two fiscal years stale.** `.../OGHIST.xlsx` serves FY25
  (data year 2023, 223 economies); `.../OGHIST_2026_07_15.xlsx` serves FY27
  (2025, 218). Both return 200 and both parse, so pointing at the undated alias
  is a silent two-year lag rather than an error. The dated file is the one
  `datacatalog.worldbank.org/search/dataset/0037712` lists, and a new vintage
  each July is a one-constant edit here — which is the whole point of the script.
* **Four year columns are formulas, not values.** FY18-FY21 hold `=IF(...)`, so
  `openpyxl` without `data_only=True` yields the *formula text* as a country's
  income group: 39 year columns become 35, and the four that vanish do so
  quietly. `data_only=True` reads the cached values Excel stored.
* **The workbook's own `Parameters` sheet is stale relative to its data.** It
  describes FY21 thresholds ("Fiscal Yeay", the publisher's typo, is the label)
  while the history sheet runs to FY27. The header row of the history sheet is
  the only authority for what the file covers; nothing else in it is.

DuckDB's `read_xlsx` is not usable here and that is worth recording next to the
`openpyxl` import: the sheet opens with three merged title rows and a blank row,
and the reader stops there, returning 3 rows of 1 column rather than failing.
`modern_data_stack.workbook` is built on it, so it is the wrong tool for this
shape despite being the repo's own.
"""

from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path
from urllib.request import urlopen

from modern_data_stack.paths import project_root

# The dated file, deliberately. See the module docstring: the undated alias
# beside it is two fiscal years behind and says so nowhere.
OGHIST_XLSX_URL = (
    "https://datacatalogfiles.worldbank.org/ddh-published/0037712/DR0095334/OGHIST_2026_07_15.xlsx"
)

SEED_DIR = project_root() / "dbt" / "seeds"
SHEET = "Country Analytical History"

# Where the header and the body sit. Row indices are 0-based into the sheet as
# read; they are constants rather than a search because the layout is a
# publisher's fixed template, and a search for "the row that looks like years"
# would silently re-anchor if the template changed rather than stopping.
FISCAL_YEAR_ROW = 4
DATA_YEAR_ROW = 5
FIRST_COUNTRY_ROW = 11

# The publisher's codes. `..` is "not classified that year" and is dropped rather
# than carried as a group: an economy the World Bank had not yet rated is a
# different fact from one it rated as low income, and a row here would make the
# two indistinguishable downstream.
CODE_TO_GROUP = {
    "L": "Low income",
    "LM": "Lower middle income",
    "UM": "Upper middle income",
    "H": "High income",
}
UNCLASSIFIED = ".."

# `LM*` appears on exactly one economy in two years — Yemen, 1987 and 1988, when
# the Yemen Arab Republic and the PDR were still separate and the workbook
# footnotes the pair. It is a lower-middle-income classification carrying a
# footnote, not a fifth group, so it maps to `LM` and the star is dropped. Left
# unmapped it would fail the guard below, which is the intended behaviour for any
# *new* starred code: stop and make a person read the footnote.
CODE_ALIASES = {"LM*": "LM"}


def parse_history(xlsx_path: Path) -> list[dict]:
    """One row per (economy, data year) the workbook classifies."""
    try:
        import openpyxl  # ty: ignore[unresolved-import]  # optional; see below
    except ModuleNotFoundError:
        # Same treatment as `scripts/build_cbam_seeds.py`: openpyxl is a
        # transcription-time dependency for a script that runs once a year, not
        # a runtime one, so it is not in any dependency group.
        sys.exit(
            "openpyxl is required: "
            "uv run --with openpyxl python -m scripts.build_income_classification_seed"
        )

    workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if SHEET not in workbook.sheetnames:
        sys.exit(f"no {SHEET!r} sheet in {xlsx_path.name}; sheets are {workbook.sheetnames}")
    rows = [list(row) for row in workbook[SHEET].iter_rows(values_only=True)]

    data_years = rows[DATA_YEAR_ROW]
    fiscal_years = rows[FISCAL_YEAR_ROW]
    year_columns = [i for i, value in enumerate(data_years) if isinstance(value, int)]
    if not year_columns:
        sys.exit(
            f"no year columns in row {DATA_YEAR_ROW + 1} of {SHEET!r}. If they read as "
            "'=IF(...' the workbook was opened without data_only=True"
        )

    out: list[dict] = []
    unknown: dict[str, str] = {}
    for row in rows[FIRST_COUNTRY_ROW:]:
        iso3 = str(row[0]).strip() if row[0] else ""
        if len(iso3) != 3:
            continue
        for column in year_columns:
            raw = row[column]
            code = str(raw).strip() if raw is not None else ""
            if not code or code == UNCLASSIFIED:
                continue
            code = CODE_ALIASES.get(code, code)
            if code not in CODE_TO_GROUP:
                unknown[code] = iso3
                continue
            out.append(
                {
                    "country_iso3": iso3,
                    "data_year": data_years[column],
                    "fiscal_year": fiscal_years[column],
                    "income_group_code": code,
                    "income_group": CODE_TO_GROUP[code],
                }
            )

    if unknown:
        # Loud rather than skipped, for `SHEET_TO_ISO3`'s reason in the CBAM
        # script: a code nobody has seen is a classification decision, and
        # dropping it writes a country out of a year it was in.
        listed = ", ".join(f"{code!r} (e.g. {iso3})" for code, iso3 in sorted(unknown.items()))
        sys.exit(f"unmapped classification codes: {listed}")

    return sorted(out, key=lambda r: (r["country_iso3"], r["data_year"]))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, help="local copy of OGHIST.xlsx")
    args = parser.parse_args()

    xlsx_path = args.xlsx
    if xlsx_path is None:
        xlsx_path = project_root() / "data" / "wb_income_classification.xlsx"
        xlsx_path.parent.mkdir(parents=True, exist_ok=True)
        print(f"downloading {OGHIST_XLSX_URL}")
        with urlopen(OGHIST_XLSX_URL) as response:  # https, pinned above
            xlsx_path.write_bytes(response.read())

    rows = parse_history(xlsx_path)
    path = SEED_DIR / "wb_income_classification.csv"
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "country_iso3",
                "data_year",
                "fiscal_year",
                "income_group_code",
                "income_group",
            ],
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(rows)

    years = sorted({r["data_year"] for r in rows})
    print(
        f"wrote {path.relative_to(project_root())}: {len(rows):,} rows, "
        f"{len({r['country_iso3'] for r in rows})} economies, {years[0]}-{years[-1]}"
    )


if __name__ == "__main__":
    main()
